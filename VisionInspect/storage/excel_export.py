"""Excel 报表导出模块.

生成"模型评估报告.xlsx"，包含四个 Sheet：
    Sheet1: 图片统计
    Sheet2: 零件统计
    Sheet3: 详细审核记录
    Sheet4: 模型信息
"""

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from data.project_data import ProjectData
from utils.common import ReviewType
from utils.logger import get_logger

logger = get_logger(__name__)

# 样式常量
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
TITLE_FONT = Font(name="微软雅黑", size=14, bold=True)
SECTION_FONT = Font(name="微软雅黑", size=12, bold=True)
_THIN_SIDE = Side(style="thin")
THIN_BORDER = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)


class ExcelExporter:
    """Excel 报告导出器."""

    # ------------------------------------------------------------------
    # 主入口
    # ------------------------------------------------------------------
    @staticmethod
    def export_report(project: ProjectData, output_dir: str = "") -> str:
        """导出完整模型评估报告.

        Args:
            project: 项目数据
            output_dir: 输出目录（默认项目根目录）

        Returns:
            导出文件路径
        """
        wb = Workbook()

        ExcelExporter._export_image_sheet(wb, project)
        ExcelExporter._export_part_sheet(wb, project)
        ExcelExporter._export_detail_sheet(wb, project)
        ExcelExporter._export_model_sheet(wb, project)

        # 确定输出路径
        if not output_dir:
            output_dir = project.root_path or "."
        report_name = "模型评估报告.xlsx"
        file_path = os.path.join(output_dir, report_name)

        # 避免覆盖：如存在则加时间戳
        if os.path.exists(file_path):
            stem = os.path.splitext(file_path)[0]
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_path = f"{stem}_{ts}.xlsx"

        wb.save(file_path)
        logger.info("Excel 报告已导出: %s", file_path)
        return file_path

    # ------------------------------------------------------------------
    # Sheet1: 图片统计
    # ------------------------------------------------------------------
    @staticmethod
    def _export_image_sheet(wb: Workbook, project: ProjectData):
        """Sheet1 图片级统计（左侧详细表格 + 右侧图片级指标统计汇总）."""
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet("图片统计")
        else:
            ws.title = "图片统计"

        ws["A1"] = "VisionInspect AI模型评估报告 - 图片级统计"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A2"].font = BODY_FONT

        headers = [
            "零件号",
            "图片文件名",
            "图片宽",
            "图片高",
            "标注数量",
            "正确",
            "漏报",
            "错报",
            "误报",
            "过杀",
            "未审核",
            "图片审核状态",
        ]
        row = 4
        ExcelExporter._write_row(ws, row, headers, header=True)
        row += 1

        # 汇总统计（图片级指标）
        total_images = 0
        reviewed_images = 0
        ann_total = 0
        summary_counts = {
            "correct": 0,
            "miss": 0,
            "wrong_class": 0,
            "false_positive": 0,
            "overkill": 0,
            "unreview": 0,
        }
        by_class = {}

        for part in project.parts:
            for img in part.images:
                total_images += 1
                if img.is_review_finished():
                    reviewed_images += 1
                counts = ExcelExporter._count_annotations(img.annotations)
                ann_total += len(img.annotations)
                for key in summary_counts:
                    summary_counts[key] += counts[key]
                for ann in img.annotations:
                    cls_name = ann.class_name or "未知"
                    if cls_name not in by_class:
                        by_class[cls_name] = {
                            "correct": 0,
                            "miss": 0,
                            "wrong_class": 0,
                            "false_positive": 0,
                            "overkill": 0,
                            "unreview": 0,
                        }
                    by_class[cls_name][ann.review_type.value] += 1

                status_text = ExcelExporter._status_text(img)
                values = [
                    part.part_id,
                    img.file_name,
                    img.width,
                    img.height,
                    len(img.annotations),
                    counts["correct"],
                    counts["miss"],
                    counts["wrong_class"],
                    counts["false_positive"],
                    counts["overkill"],
                    counts["unreview"],
                    status_text,
                ]
                ExcelExporter._write_row(ws, row, values)
                row += 1

        ExcelExporter._auto_width(ws, headers)

        # 右侧：图片级指标统计汇总（M 列空出，从 N 列起）
        col1 = 14  # N
        ExcelExporter._write_summary_title(ws, 4, col1, "图片级指标统计汇总")

        summary_headers = ["指标", "数值"]
        ExcelExporter._write_row(ws, 5, summary_headers, header=True, start_col=col1)

        summary_items = [
            ("总图片数", total_images),
            ("已审核图片数", reviewed_images),
            ("未审核图片数", total_images - reviewed_images),
            ("标注总数", ann_total),
            ("正确", summary_counts["correct"]),
            ("漏报", summary_counts["miss"]),
            ("错报", summary_counts["wrong_class"]),
            ("误报", summary_counts["false_positive"]),
            ("过杀", summary_counts["overkill"]),
            ("未审核标注", summary_counts["unreview"]),
        ]
        data_row = 6
        for name, value in summary_items:
            ExcelExporter._write_row(ws, data_row, [name, value], start_col=col1)
            data_row += 1

        # 缺陷类别分布子表
        data_row += 1
        ExcelExporter._write_summary_title(ws, data_row, col1, "缺陷类别分布")
        data_row += 1
        class_headers = ["类别", "正确", "漏报", "错报", "误报", "过杀", "未审"]
        ExcelExporter._write_row(ws, data_row, class_headers, header=True, start_col=col1)
        data_row += 1
        for cls_name in sorted(by_class.keys()):
            c = by_class[cls_name]
            ExcelExporter._write_row(
                ws,
                data_row,
                [
                    cls_name,
                    c["correct"],
                    c["miss"],
                    c["wrong_class"],
                    c["false_positive"],
                    c["overkill"],
                    c["unreview"],
                ],
                start_col=col1,
            )
            data_row += 1

        # 汇总区列宽
        for col_idx in range(col1, col1 + len(class_headers)):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 14

    # ------------------------------------------------------------------
    # Sheet2: 零件统计
    # ------------------------------------------------------------------
    @staticmethod
    def _export_part_sheet(wb: Workbook, project: ProjectData):
        """Sheet2 零件级统计."""
        ws = wb.create_sheet("零件统计")

        ws["A1"] = "VisionInspect AI模型评估报告 - 零件级统计"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A2"].font = BODY_FONT

        # 第一部分：缺陷实例统计
        headers = [
            "零件号",
            "缺陷实例ID",
            "缺陷类别",
            "关联标注数",
            "正确数",
            "漏报数",
            "错报数",
            "误报数",
            "过杀数",
            "零件级结果",
        ]
        row = 4
        ExcelExporter._write_row(ws, row, headers, header=True)
        row += 1

        for part in project.parts:
            # 1. 已绑定的缺陷实例
            for defect in part.defects:
                result = defect.get_part_result()
                values = [
                    part.part_id,
                    defect.defect_id,
                    defect.class_name,
                    len(defect.annotation_bindings),
                    result["correct_count"],
                    result["miss_count"],
                    result["wrong_count"],
                    result["fp_count"],
                    result["overkill_count"],
                    ExcelExporter._part_result_text(result["status"]),
                ]
                ExcelExporter._write_row(ws, row, values)
                row += 1

            # 2. 未绑定缺陷实例的过杀检测框（每框独立列出）
            bound_ann_ids = set()
            for defect in part.defects:
                for binding in defect.annotation_bindings:
                    bound_ann_ids.add((binding.get("image_id"), binding.get("annotation_id")))

            unbound_overkill = []
            for img in part.images:
                for ann in img.annotations:
                    key = (img.image_id, ann.id)
                    if key in bound_ann_ids:
                        continue
                    if ann.review_type == ReviewType.OVERKILL:
                        unbound_overkill.append((img, ann))

            if unbound_overkill:
                ws.merge_cells(
                    start_row=row,
                    start_column=1,
                    end_row=row,
                    end_column=len(headers),
                )
                title_cell = ws.cell(
                    row=row, column=1, value="—— 未绑定缺陷实例的过杀检测框 ——"
                )
                title_cell.font = SECTION_FONT
                title_cell.alignment = Alignment(horizontal="center", vertical="center")
                row += 1
                for img, ann in unbound_overkill:
                    values = [
                        part.part_id,
                        "未绑定",
                        ann.review_class or ann.class_name,
                        1,
                        0,
                        0,
                        0,
                        0,
                        1,
                        ExcelExporter._part_result_text("overkill"),
                    ]
                    ExcelExporter._write_row(ws, row, values)
                    row += 1

        # 第二部分：零件级指标汇总
        row += 2
        ws.cell(row=row, column=1, value="零件级指标汇总").font = Font(
            name="微软雅黑", size=12, bold=True
        )
        row += 1

        part_stats = ExcelExporter._compute_part_stats(project)
        summary_headers = ["指标", "数值"]
        ExcelExporter._write_row(ws, row, summary_headers, header=True)
        row += 1

        summary_items = [
            ("零件总数", part_stats["total_parts"]),
            ("缺陷实例总数", part_stats["total_defects"]),
            ("检测成功数", part_stats["detected"]),
            ("漏检数", part_stats["missed"]),
            ("误检数", part_stats["false_positive"]),
            ("过杀数", part_stats.get("overkill", 0)),
            ("Precision", f"{part_stats['precision']:.4f}" if part_stats["precision"] is not None else "-"),
            ("Recall", f"{part_stats['recall']:.4f}" if part_stats["recall"] is not None else "-"),
            ("F1 Score", f"{part_stats['f1']:.4f}" if part_stats["f1"] is not None else "-"),
        ]
        for name, value in summary_items:
            ExcelExporter._write_row(ws, row, [name, value])
            row += 1

        ExcelExporter._auto_width(ws, headers)

    # ------------------------------------------------------------------
    # Sheet3: 详细审核记录
    # ------------------------------------------------------------------
    @staticmethod
    def _export_detail_sheet(wb: Workbook, project: ProjectData):
        """Sheet3 详细审核记录."""
        ws = wb.create_sheet("详细审核记录")

        ws["A1"] = "VisionInspect AI模型评估报告 - 详细审核记录"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A2"].font = BODY_FONT

        headers = [
            "零件号",
            "图片文件名",
            "标注ID",
            "类别名称",
            "置信度",
            "来源",
            "审核类型",
            "审核类别",
            "检测框(x1,y1,x2,y2)",
            "关联缺陷ID",
            "审核人员",
            "审核时间",
        ]
        row = 4
        ExcelExporter._write_row(ws, row, headers, header=True)
        row += 1

        for part in project.parts:
            for img in part.images:
                for ann in img.annotations:
                    bbox_str = f"({ann.bbox[0]},{ann.bbox[1]},{ann.bbox[2]},{ann.bbox[3]})"
                    values = [
                        part.part_id,
                        img.file_name,
                        ann.id,
                        ann.class_name,
                        round(ann.confidence, 4),
                        ann.source.value,
                        ann.review_type.display_name,
                        ann.review_class,
                        bbox_str,
                        ann.defect_id or "",
                        img.review_user,
                        img.review_time,
                    ]
                    ExcelExporter._write_row(ws, row, values)
                    row += 1

        ExcelExporter._auto_width(ws, headers)

    # ------------------------------------------------------------------
    # Sheet4: 模型信息
    # ------------------------------------------------------------------
    @staticmethod
    def _export_model_sheet(wb: Workbook, project: ProjectData):
        """Sheet4 模型信息汇总."""
        ws = wb.create_sheet("模型信息")

        ws["A1"] = "VisionInspect AI模型评估报告 - 模型信息"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A2"].font = BODY_FONT

        headers = ["项目名称", "数据根目录", "模型名称", "模型版本", "模型文件", "推理时间", "审核人员"]
        row = 4
        ExcelExporter._write_row(ws, row, headers, header=True)
        row += 1

        # 从图片 JSON 中收集模型信息（可能有多个模型版本）
        model_entries = {}
        for part in project.parts:
            for img in part.images:
                key = (img.model_name, img.model_version, img.model_file)
                if key not in model_entries:
                    model_entries[key] = {
                        "model_name": img.model_name,
                        "model_version": img.model_version,
                        "model_file": img.model_file,
                        "infer_time": img.infer_time,
                        "review_user": img.review_user,
                        "image_count": 0,
                    }
                model_entries[key]["image_count"] += 1

        if not model_entries:
            values = [project.project_name, project.root_path, "", "", "", "", ""]
            ExcelExporter._write_row(ws, row, values)
        else:
            for key, entry in model_entries.items():
                values = [
                    project.project_name,
                    project.root_path,
                    entry["model_name"],
                    entry["model_version"],
                    entry["model_file"],
                    entry["infer_time"],
                    entry["review_user"],
                ]
                ExcelExporter._write_row(ws, row, values)
                row += 1

        # 类别信息
        row += 2
        ws.cell(row=row, column=1, value="模型识别类别").font = Font(name="微软雅黑", size=12, bold=True)
        row += 1
        class_headers = ["序号", "类别名称"]
        ExcelExporter._write_row(ws, row, class_headers, header=True)
        row += 1
        if project.parts:
            part0 = project.parts[0]
            if part0.images:
                classes = set()
                for img in part0.images:
                    for ann in img.annotations:
                        if ann.class_name:
                            classes.add(ann.class_name)
                for idx, cls_name in enumerate(sorted(classes), 1):
                    ExcelExporter._write_row(ws, row, [idx, cls_name])
                    row += 1

        ExcelExporter._auto_width(ws, headers)

    # ------------------------------------------------------------------
    # 工具方法
    # ------------------------------------------------------------------
    @staticmethod
    def _count_annotations(annotations) -> dict:
        """统计标注审核类型数量."""
        counts = {
            "correct": 0,
            "miss": 0,
            "wrong_class": 0,
            "false_positive": 0,
            "overkill": 0,
            "unreview": 0,
        }
        for ann in annotations:
            key = ann.review_type.value
            if key in counts:
                counts[key] += 1
        return counts

    @staticmethod
    def _status_text(img) -> str:
        """图片审核状态中文文本."""
        status_map = {
            "unreviewed": "未审核",
            "partial": "部分审核",
            "finished": "审核完成",
            "has_defect": "存在缺陷",
        }
        text = status_map.get(img.review_status_text)
        return text if text else str(img.review_status_text)

    @staticmethod
    def _part_result_text(status: str) -> str:
        """零件级结果中文文本."""
        status_map = {
            "detected": "检测成功",
            "missed": "漏检",
            "false_positive": "误检",
            "wrong_class": "错检",
            "overkill": "过杀",
            "unknown": "未知",
        }
        return status_map.get(status, status)

    @staticmethod
    def _compute_part_stats(project: ProjectData) -> dict:
        """计算零件级汇总指标.

        统计口径与 DefectManager.get_part_result() 一致：
            - 已绑定的缺陷实例：每个实例 = 1 个缺陷，按实例最终状态计数；
              过杀数按实例内所有"过杀"检测框数量累计（不被实例状态吞掉）。
            - 未绑定的检测框：每框独立 = 1 个缺陷，按审核类型计数。

        Recall = TP / (TP + FN)，Precision = TP / (TP + FP)，F1 = 2PR/(P+R)
        其中 TP=检测成功，FP=误检+错检，FN=漏检。
        """
        stats = {
            "total_parts": len(project.parts),
            "total_defects": 0,
            "detected": 0,
            "missed": 0,
            "false_positive": 0,
            "wrong_class": 0,
            "overkill": 0,
            "precision": None,
            "recall": None,
            "f1": None,
        }

        for part in project.parts:
            # 1. 已绑定的缺陷实例
            for defect in part.defects:
                result = defect.get_part_result()
                stats["total_defects"] += 1
                status = result["status"]
                if status == "detected":
                    stats["detected"] += 1
                elif status == "missed":
                    stats["missed"] += 1
                elif status == "false_positive":
                    stats["false_positive"] += 1
                elif status == "wrong_class":
                    stats["wrong_class"] += 1
                stats["overkill"] += result.get("overkill_count", 0)

            # 2. 未绑定的检测框（每框独立计数）
            bound_ann_ids = set()
            for defect in part.defects:
                for binding in defect.annotation_bindings:
                    bound_ann_ids.add((binding.get("image_id"), binding.get("annotation_id")))

            for img in part.images:
                for ann in img.annotations:
                    key = (img.image_id, ann.id)
                    if key in bound_ann_ids:
                        continue  # 已绑定到缺陷实例，跳过
                    stats["total_defects"] += 1
                    rt = ann.review_type.value
                    if rt == "correct":
                        stats["detected"] += 1
                    elif rt == "miss":
                        stats["missed"] += 1
                    elif rt == "false_positive":
                        stats["false_positive"] += 1
                    elif rt == "wrong_class":
                        stats["wrong_class"] += 1
                    elif rt == "overkill":
                        stats["overkill"] += 1
                    # unreview 不参与零件级结果统计

        # Precision = TP / (TP + FP)
        tp = stats["detected"]
        fp = stats["false_positive"] + stats["wrong_class"]
        fn = stats["missed"]

        if tp + fp > 0:
            stats["precision"] = tp / (tp + fp)
        if tp + fn > 0:
            stats["recall"] = tp / (tp + fn)
        p = stats["precision"]
        r = stats["recall"]
        if p is not None and r is not None and p + r > 0:
            stats["f1"] = 2 * p * r / (p + r)

        return stats

    @staticmethod
    def _write_summary_title(ws, row: int, col: int, text: str):
        """写入汇总区小节标题（跨两列居中）."""
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = SECTION_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def _write_row(ws, row: int, values: list, header: bool = False, start_col: int = 1):
        """写一行数据并应用样式."""
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=start_col + col_idx - 1, value=value)
            cell.font = HEADER_FONT if header else BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if header:
                cell.fill = HEADER_FILL

    @staticmethod
    def _auto_width(ws, headers: list, min_width: int = 10, max_width: int = 40):
        """自动调整列宽."""
        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = len(headers[col_idx - 1])
            for cell in ws[col_letter]:
                if cell.value is not None:
                    text_len = len(str(cell.value))
                    if text_len > max_len:
                        max_len = text_len
            width = min(max(max_len * 1.2 + 4, min_width), max_width)
            ws.column_dimensions[col_letter].width = width